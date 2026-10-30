"""Regenerate the charts on the Data page.

Run from the site folder:  python3 scripts/make_data_figures.py
Reads the Excel files in public/data/ and writes:
  - src/data/charts/*.json  (data for the interactive charts on the page)
  - public/data/fig/*.svg   (static versions, e.g. for slides or social posts)
"""
import datetime as dt
import json
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates

SURF = "#ffffff"; INK2 = "#52514e"; MUTED = "#8a8984"; GRID = "#ebeae6"; BLUE = "#2a78d6"
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "svg.fonttype": "none"})


def chart(dates, values, width_days, ylabel, out, xmin, xmax, step=5):
    fig, ax = plt.subplots(figsize=(8.2, 2.1), dpi=100, facecolor=SURF)
    ax.set_facecolor(SURF)
    ax.bar(dates, values, width=width_days, color=BLUE, linewidth=0)
    ax.axhline(0, color=MUTED, linewidth=0.8)
    for s in ax.spines.values():
        s.set_visible(False)
    ax.grid(axis="y", color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)
    ax.tick_params(colors=INK2, length=0, labelsize=12)
    ax.xaxis.set_major_locator(mdates.YearLocator(step))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    ax.set_xlim(xmin, xmax)
    if ylabel:
        ax.set_ylabel(ylabel, color=INK2, fontsize=11)
    fig.tight_layout(pad=0.4)
    fig.savefig(out, facecolor=SURF)
    plt.close(fig)


# Monetary policy instrument (MPI), monthly
ws = openpyxl.load_workbook("public/data/MAR_IVs.xlsx", data_only=True).active
rows = [r for r in list(ws.iter_rows(values_only=True))[1:] if r[0] is not None]
hdr = [c for c in next(ws.iter_rows(values_only=True))]
ci = hdr.index("MPI")
json.dump([{"t": f"{int(r[0])}M{int(r[1])}", "x": int(r[0]) + (int(r[1]) - 0.5) / 12, "v": round(r[ci], 6)} for r in rows],
          open("src/data/charts/monetary_policy.json", "w"))
chart([dt.date(int(r[0]), int(r[1]), 15) for r in rows], [r[ci] for r in rows], 25, "",
      "public/data/fig/monetary_policy.svg", dt.date(1990, 7, 1), dt.date(2021, 7, 1))

# Fiscal: Ramey forecast error (mean), quarterly
ws = openpyxl.load_workbook("public/data/fiscal_news_proxies.xlsx", data_only=True)["Ramey"]
rows = list(ws.iter_rows(values_only=True))
ci = rows[0].index("FE (mean)")
def q2d(s):
    y, q = s.split("Q"); return dt.date(int(y), 3 * int(q) - 1, 15)
data = [(q2d(r[0]), r[ci]) for r in rows[1:] if r[0] and r[ci] is not None]
json.dump([{"t": r[0], "x": int(r[0][:4]) + (int(r[0][-1]) - 0.5) / 4, "v": round(r[ci], 4)}
           for r in rows[1:] if r[0] and r[ci] is not None],
          open("src/data/charts/fiscal_news.json", "w"))
chart([d for d, _ in data], [v for _, v in data], 70, "annualized %",
      "public/data/fig/fiscal_news.svg", dt.date(1981, 1, 1), dt.date(2027, 1, 1))
print("figures written")
